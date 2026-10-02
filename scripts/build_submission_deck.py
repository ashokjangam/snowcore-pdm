"""Build the six-slide PNEUMORA deck from the official hackathon template."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEAM_URL = (
    "https://hack2skill.com/event/cococlihack-gccedition/dashboard/"
    "team-management?utm_source=hack2skill&utm_medium=homepage"
)
LEADER_URL = "https://www.linkedin.com/in/jangam-ashok-53a95812b/"
APP_URL = "https://app.snowflake.com/streamlit/FMXJOWH/BRC04642/#/apps/df4d72ofa3zbhz6wswdh"
REPO_URL = "https://github.com/ashokjangam/snowcore-pdm/tree/feat/pneumora-metropt"

WHITE = "F4FAFD"
MUTED = "A5B7C3"
CYAN = "36A9E1"
TEAL = "31D7C5"
AMBER = "F2B84B"
COPPER = "D9844F"
RED = "EF6A72"
BLUE = "0A5CA8"
DARK = "0C1822"
PANEL = "122431"
GRID = "29404E"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def remove_non_picture_shapes(slide) -> None:
    for shape in list(slide.shapes):
        if shape.shape_type != 13:  # PICTURE
            shape._element.getparent().remove(shape._element)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 15,
    color: str = WHITE,
    bold: bool = False,
    font: str = "Calibri",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=15, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for spec in runs:
        run = paragraph.add_run()
        run.text = spec["text"]
        run.font.name = spec.get("font", "Calibri")
        run.font.size = Pt(spec.get("size", size))
        run.font.bold = spec.get("bold", False)
        run.font.color.rgb = rgb(spec.get("color", color))
        if spec.get("url"):
            run.hyperlink.address = spec["url"]
    return box


def add_title(slide, title: str, kicker: str) -> None:
    add_text(slide, kicker.upper(), 0.48, 0.50, 5.4, 0.22, size=9, color=TEAL, bold=True)
    add_text(slide, title, 0.46, 0.76, 8.9, 0.55, size=27, color=DARK, bold=True, font="Cambria")


def add_footer(slide, text: str) -> None:
    add_text(slide, text, 0.48, 5.31, 9.0, 0.17, size=7, color="667D8C")


def shape_box(slide, x, y, w, h, fill, *, radius=True, line=GRID):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    return shape


def circle(slide, x, y, diameter, fill, line=GRID):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1.4)
    return shape


def arrow(slide, x, y, w=0.36, h=0.24, color=CYAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def stat(slide, value, label, note, x, y, w, tone=TEAL):
    shape_box(slide, x, y, w, 1.45, PANEL)
    add_text(slide, value, x + 0.16, y + 0.12, w - 0.32, 0.42, size=26, color=tone, bold=True)
    add_text(slide, label, x + 0.16, y + 0.59, w - 0.32, 0.24, size=11, color=WHITE, bold=True)
    add_text(slide, note, x + 0.16, y + 0.88, w - 0.32, 0.42, size=9, color=MUTED)


def fill_title_slide(slide) -> None:
    remove_non_picture_shapes(slide)
    add_text(slide, "PNEUMORA", 6.72, 3.28, 2.78, 0.48, size=29, color=BLUE, bold=True, font="Cambria")
    add_text(slide, "EARLY AIR-LEAK PREDICTION ON REAL COMPRESSOR SENSORS", 6.73, 3.78, 2.74, 0.48, size=10, color=DARK, bold=True)
    add_rich_text(
        slide,
        [
            {"text": "Team Name: ", "bold": True, "color": DARK},
            {"text": "TRidents", "bold": True, "color": BLUE, "url": TEAM_URL},
        ],
        0.48, 3.44, 5.6, 0.28, size=13, color=DARK,
    )
    add_rich_text(
        slide,
        [
            {"text": "Team Leader: ", "bold": True, "color": DARK},
            {"text": "Ashok Jangam", "bold": True, "color": BLUE, "url": LEADER_URL},
        ],
        0.48, 3.82, 5.6, 0.28, size=13, color=DARK,
    )
    add_text(slide, "Team Size: 4", 0.48, 4.20, 5.6, 0.28, size=13, color=DARK, bold=True)
    add_text(
        slide,
        "Problem Statement: Predictive Maintenance and OEE Command Center",
        0.48, 4.58, 8.95, 0.30, size=13, color=DARK, bold=True,
    )
    add_text(slide, "Nine real failures. Three compressors. A breakdown predicted hours ahead.",
             0.48, 5.05, 8.95, 0.26, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


def fill_problem_slide(slide) -> None:
    remove_non_picture_shapes(slide)
    add_title(slide, "A leaking compressor takes a train out of service.", "01 · Problem brief")
    circle(slide, 0.65, 1.55, 2.25, "102C3A", TEAL)
    add_text(slide, "REAL", 1.17, 1.86, 1.2, 0.34, size=22, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "MetroPT sensors\npressure · current\noil temp · load", 0.94, 2.30, 1.68, 0.84,
             size=11, color=WHITE, align=PP_ALIGN.CENTER)
    circle(slide, 7.08, 1.55, 2.25, "2E2513", AMBER)
    add_text(slide, "LABELLED", 7.38, 1.86, 1.65, 0.34, size=20, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "synthetic work orders\nparts · crew\nexample-factory OEE", 7.34, 2.30, 1.75, 0.84,
             size=11, color=WHITE, align=PP_ALIGN.CENTER)
    shape_box(slide, 3.48, 1.70, 3.05, 1.90, "1A1D22", radius=False, line=RED)
    add_text(slide, "EXISTING ALARM", 3.62, 1.86, 2.77, 0.30, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "3 of 7", 3.62, 2.22, 2.77, 0.50, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "air leaks flagged \u2265 2 h before removal\n0.49 false alerts per healthy day", 3.62, 2.84, 2.77, 0.55,
             size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Same asset, one thread", 0.75, 4.08, 1.6, 0.42, size=11, color=TEAL, bold=True)
    add_text(
        slide,
        "PNEUMORA reads the compressor\u2019s real sensor streams, predicts 6 of 7 air-leak breakdowns hours ahead with about a "
        "ninth of the alarm\u2019s false alerts, and turns each one into a cited, recorded maintenance decision.",
        2.30, 3.96, 7.05, 0.80, size=15, color=DARK, bold=True,
    )
    add_footer(slide, "Sources: MetroPT-3 DOI 10.24432/C5VW3R · MetroPT 2022 Zenodo 6854240 · MetroPT-2 Zenodo 7766691 · 9 documented failures")


def fill_architecture_slide(slide) -> None:
    remove_non_picture_shapes(slide)
    add_title(slide, "Raw sensor cycles in. A recorded decision out.", "02 · Architecture")
    shape_box(slide, 0.52, 1.50, 1.65, 1.12, "102C3A")
    add_text(slide, "MetroPT raw", 0.62, 1.72, 1.45, 0.25, size=14, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "1 Hz / 10 s · 3 units", 0.62, 2.09, 1.45, 0.22, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(slide, 2.30, 1.94)
    shape_box(slide, 2.82, 1.36, 2.15, 1.42, "122431")
    add_text(slide, "CORE + ML", 3.05, 1.62, 1.68, 0.26, size=14, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "cycle physics · 5-min telemetry\nANOMALY_DETECTION", 2.95, 2.00, 1.90, 0.50, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    shape_box(slide, 0.52, 3.42, 1.65, 1.12, "2E2513")
    add_text(slide, "Synthetic CMMS", 0.58, 3.66, 1.53, 0.25, size=13, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "labelled, same asset", 0.62, 4.03, 1.45, 0.22, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(slide, 2.30, 3.86, color=AMBER)
    shape_box(slide, 2.82, 3.28, 2.15, 1.42, "282316")
    add_text(slide, "OPS", 3.05, 3.54, 1.68, 0.26, size=14, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "work orders · parts · crew\nACTION_LOG", 2.95, 3.92, 1.90, 0.50, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    arrow(slide, 5.16, 2.38, 0.42, 0.26)
    arrow(slide, 5.16, 3.28, 0.42, 0.26, AMBER)
    shape_box(slide, 5.76, 1.46, 3.68, 3.02, "0E2029", line=TEAL)
    add_text(slide, "PNEUMORA app", 6.00, 1.70, 3.20, 0.33, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for x, label, note, tone in [
        (6.02, "PREDICT", "leak physics +\nnative ML", CYAN),
        (7.10, "EXPLAIN", "AI_COMPLETE,\ntyped + cited", TEAL),
        (8.18, "ACT", "idempotent\nMERGE", AMBER),
    ]:
        circle(slide, x, 2.42, 0.88, PANEL, tone)
        add_text(slide, label, x + 0.01, 2.76, 0.86, 0.16, size=8, color=tone, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, note, x - 0.08, 3.38, 1.04, 0.40, size=8, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "PNEUMORA.OPS.RECORD_ACTION", 6.10, 4.05, 3.00, 0.20,
             size=8, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Snowflake: warehouse-runtime Streamlit · SNOWFLAKE.ML.ANOMALY_DETECTION · Cortex AI_COMPLETE · SQL procedure · isolated PNEUMORA database")


def fill_evidence_slide(slide) -> None:
    remove_non_picture_shapes(slide)
    add_title(slide, "It predicts the breakdown hours ahead. Promoted.", "03 · Measured results")
    stat(slide, "6 of 7", "Air leaks predicted \u2265 2 h ahead", "Held out by compressor · existing alarm 3 of 7 · 0.056 vs 0.49 false alerts/day", 0.55, 1.48, 2.75, TEAL)
    stat(slide, "5.8 h", "Median warning before removal", "PROMOTED_CROSS_VALIDATED · all five pre-declared gates passed · chance p = 3e-7", 3.62, 1.48, 2.75, CYAN)
    stat(slide, "1 of 9", "Warned before the leak began", "Leaks start abruptly · best of six approaches · it predicts the breakdown, not the leak", 6.69, 1.48, 2.75, AMBER)
    add_text(slide, "Snowflake native ML + gradual leaks", 0.62, 3.27, 3.00, 0.28, size=13, color=DARK, bold=True)
    add_text(slide, "ANOMALY_DETECTION: 4 of 4 in time. Injected 2\u00d7 air-loss leak caught 28 of 30 (control 7).", 0.62, 3.66, 2.80, 0.82, size=11, color="087F72", bold=True)
    add_text(slide, "What we generated", 3.62, 3.27, 2.76, 0.28, size=13, color=DARK, bold=True)
    add_text(slide, "Replayed onsets and physics-injected leaks, all SYNTHETIC_TRAINING_ONLY, never used to score.", 3.62, 3.66, 2.76, 0.82, size=11, color="B07400", bold=True)
    add_text(slide, "Caveats we show on screen", 6.70, 3.27, 2.68, 0.28, size=13, color=DARK, bold=True)
    add_text(slide, "Protocol v2 written after v1 failed on false alerts; no untouched data left. No oil leaks, RUL or savings.", 6.70, 3.66, 2.68, 0.82, size=11, color=RED, bold=True)
    add_footer(slide, "Predicted = alert from 2 h before leak start to 2 h before reported end (dataset owners\u2019 protocol) · settings chosen on the other two compressors")


def fill_demo_slide(slide) -> None:
    remove_non_picture_shapes(slide)
    add_title(slide, "The 4-minute proof: predict \u2192 explain \u2192 act", "04 · End-to-end CoCo CLI workflow")
    steps = [
        ("1", "PREDICT", "CoCo reads failures, alerts and native-ML results. The app shows the held-out prediction record.", CYAN),
        ("2", "EXPLAIN", "Observed F04 facts go to Cortex. Every number cites its field; no invented parts.", TEAL),
        ("3", "ACT", "The crew opens an inspection. A Snowflake MERGE writes one reload-safe row.", AMBER),
    ]
    for index, (number, title, body, tone) in enumerate(steps):
        x = 0.62 + index * 3.10
        circle(slide, x, 1.66, 0.68, tone, tone)
        add_text(slide, number, x, 1.82, 0.68, 0.24, size=17, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.86, 1.72, 1.70, 0.28, size=17, color=tone, bold=True)
        add_text(slide, body, x, 2.60, 2.57, 1.00, size=12, color=DARK)
        if index < 2:
            arrow(slide, x + 2.61, 2.02, 0.32, 0.22, tone)
    shape_box(slide, 1.05, 4.16, 7.90, 0.62, "151F26", line=GRID)
    add_text(slide, "BLACK-BOX RESULT", 1.28, 4.34, 1.42, 0.18, size=9, color=TEAL, bold=True)
    add_text(slide, "same decision \u2192 same action id \u2192 second call deduplicated=true", 2.78, 4.28, 5.75, 0.28,
             size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Saved prompts: docs/demo/01_detect.txt · 02_rca.txt · 03_action.txt")


def fill_thank_you(slide) -> None:
    remove_non_picture_shapes(slide)
    shape_box(slide, 0.56, 4.66, 8.88, 0.54, "07131E", line=GRID)
    add_rich_text(
        slide,
        [
            {"text": "LIVE  ", "bold": True, "color": TEAL},
            {"text": "PNEUMORA", "bold": True, "color": WHITE, "url": APP_URL},
            {"text": "     CODE  ", "bold": True, "color": CYAN},
            {"text": "github.com/ashokjangam/snowcore-pdm", "color": WHITE, "url": REPO_URL},
        ],
        0.76, 4.82, 8.46, 0.23, size=10, color=WHITE, align=PP_ALIGN.CENTER,
    )
    add_text(slide, "Predict the breakdown. Explain with evidence. Record the decision. Say where prediction ends.",
             1.20, 5.22, 7.60, 0.22, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def build(template: Path, output: Path) -> None:
    presentation = Presentation(str(template))
    if len(presentation.slides) != 6:
        raise ValueError(f"Expected the official six-slide template, found {len(presentation.slides)} slides")
    fill_title_slide(presentation.slides[0])
    fill_problem_slide(presentation.slides[1])
    fill_architecture_slide(presentation.slides[2])
    fill_evidence_slide(presentation.slides[3])
    fill_demo_slide(presentation.slides[4])
    fill_thank_you(presentation.slides[5])
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.template, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
